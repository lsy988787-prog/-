"""由课程 JSON 生成给老师用的审校表。"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
course = json.load(open(src, encoding="utf-8"))

FONT = "Arial"
f_body = Font(name=FONT, size=10)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=14, bold=True)
f_bold = Font(name=FONT, size=10, bold=True)
fill_head = PatternFill("solid", start_color="1F4E78")
fill_input = PatternFill("solid", start_color="FFF2CC")
fill_prio = PatternFill("solid", start_color="FCE4D6")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

PRIORITY = {"per-unit": "标幺制", "per-unit-network": "标幺制", "practical-3ph-calculation": "三相短路",
            "infinite-bus-short-circuit": "三相短路", "generator-short-circuit": "三相短路",
            "single-line-ground-fault": "单相接地", "sequence-networks": "单相接地",
            "ac-traction-system": "牵引供电", "traction-negative-sequence": "牵引供电", "dc-traction-system": "牵引供电"}

titles = {c["id"]: c["title"] for ch in course["chapters"] for c in ch["concepts"]}

wb = Workbook()

# ---------- 使用说明 ----------
ws = wb.active
ws.title = "使用说明"
lines = [
    ("《电力系统分析》知识点审校表", f_title),
    ("课程：%s（课程 ID %s，内容版本 %s，生成日期 %s）" % (course["title"], course["id"], course["contentVersion"], course["generatedOn"]), f_body),
    ("内容来源：以中国大学 MOOC 西南交通大学《电力系统分析》公开课程主题为蓝本的原创教学草稿，不含视频、课件或教材原文。程序 MIT，教学内容 CC BY 4.0。", f_body),
    ("", f_body),
    ("怎么填", f_bold),
    ("1. 在“知识点审校”表中逐行阅读内容，只需填写浅黄色列：审校结论、难度、修改意见、审校人、审校日期。", f_body),
    ("2. “审校结论”从下拉选项中选：通过 / 小改 / 需重写。选“小改”或“需重写”时请在“修改意见”中写明要改的内容。", f_body),
    ("3. 浅橙色底的行是建议优先审校的知识点（标幺制、三相短路、单相接地、牵引供电）。", f_body),
    ("4. “公式核对”表逐条列出全部公式、符号单位和适用条件，可在浅黄色列标注是否正确。", f_body),
    ("5. “审校进度”表自动统计各章已审数量和结论分布，无需填写。", f_body),
    ("", f_body),
    ("填写示例（一行）", f_bold),
    ("审校结论：小改 ｜ 难度：合适 ｜ 修改意见：例题第 3 步建议补充“近似法忽略变比差异”的说明；自测题答案保留两位小数 ｜ 审校人：张老师 ｜ 审校日期：2026-11-05", f_body),
    ("", f_body),
    ("请重点关注", f_bold),
    ("· 有名值与标幺值、三相容量与线/相电压的搭配是否一致；功率和电流方向是否交代清楚。", f_body),
    ("· dq0 变换采用幅值不变、q 轴超前 d 轴 90° 的约定，全课程是否一致。", f_body),
    ("· 各类不对称故障的序网连接与边界条件是否正确；交流牵引与直流牵引、交流电网“直流潮流”的边界是否讲清。", f_body),
    ("· 引用的国家标准限值（GB/T 12325、15945、15543、14549）是否为现行版本。", f_body),
    ("· 教学滑块范围和例题参数是教学近似，不作为工程标准或保护整定依据。", f_body),
]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 120

# ---------- 知识点审校 ----------
wr = wb.create_sheet("知识点审校")
cols = [("序号", 6), ("章节", 14), ("知识点 ID", 22), ("知识点", 20), ("优先审校", 9), ("前置知识", 22), ("简介", 36), ("讲解", 70),
        ("关键公式", 44), ("例题", 40), ("解答步骤", 50), ("例题答案", 24), ("常见误区（错误理解 → 正确理解）", 50), ("自测题", 34),
        ("自测答案与解析", 40), ("可视化（计算器与参数）", 36),
        ("审校结论", 11), ("难度", 9), ("修改意见", 44), ("审校人", 10), ("审校日期", 12)]
INPUT_FROM = 17
for j, (h, w) in enumerate(cols, 1):
    c = wr.cell(row=1, column=j, value=h)
    c.font, c.fill, c.border = f_head, fill_head, border
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    wr.column_dimensions[get_column_letter(j)].width = w
wr.row_dimensions[1].height = 30

r = 2
for ch in course["chapters"]:
    for c in ch["concepts"]:
        we = c["workedExample"]
        pr = c["practice"][0]
        viz = c["visualization"]
        viz_text = ""
        if viz:
            viz_text = viz["calculator"] + "：" + "；".join(
                "%s %s~%s %s（默认 %s）" % (p["label"], p["min"], p["max"], p["unit"], p["default"]) for p in viz["params"]) + "。假设：" + viz["assumptions"]
        vals = [
            r - 1, "%d %s" % (ch["order"], ch["title"]), c["id"], c["title"], PRIORITY.get(c["id"], ""),
            "、".join(titles[p] for p in c["prerequisites"]) or "无",
            c["summary"], c["explanation"],
            "\n".join("%d) %s%s" % (k + 1, f["display"], ("（" + f["conditions"] + "）") if f["conditions"] else "") for k, f in enumerate(c["keyFormulas"])),
            we["problem"], "\n".join("%d. %s" % (k + 1, s) for k, s in enumerate(we["steps"])), we["answer"],
            "\n".join("× %s\n√ %s" % (m["wrong"], m["correct"]) for m in c["misconceptions"]),
            pr["question"], "答案：%s\n解析：%s" % (pr["answer"], pr["explanation"]), viz_text,
            None, None, None, None, None,
        ]
        for j, v in enumerate(vals, 1):
            cell = wr.cell(row=r, column=j, value=v)
            cell.font, cell.alignment, cell.border = f_body, wrap, border
            if j >= INPUT_FROM:
                cell.fill = fill_input
            elif c["id"] in PRIORITY:
                cell.fill = fill_prio
        wr.cell(row=r, column=21).number_format = "yyyy-mm-dd"
        r += 1
last = r - 1
wr.freeze_panes = "E2"
wr.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), last)

dv1 = DataValidation(type="list", formula1='"通过,小改,需重写"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"偏易,合适,偏难"', allow_blank=True)
wr.add_data_validation(dv1)
wr.add_data_validation(dv2)
dv1.add("Q2:Q%d" % last)
dv2.add("R2:R%d" % last)

# ---------- 公式核对 ----------
wf = wb.create_sheet("公式核对")
fcols = [("序号", 6), ("知识点 ID", 22), ("知识点", 20), ("公式", 48), ("LaTeX 源", 48), ("符号与单位", 50), ("适用条件", 30), ("是否正确", 10), ("意见", 40)]
for j, (h, w) in enumerate(fcols, 1):
    c = wf.cell(row=1, column=j, value=h)
    c.font, c.fill, c.border = f_head, fill_head, border
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    wf.column_dimensions[get_column_letter(j)].width = w
r = 2
for ch in course["chapters"]:
    for c in ch["concepts"]:
        for f in c["keyFormulas"]:
            vals = [r - 1, c["id"], c["title"], f["display"], f["latex"],
                    "\n".join("%s：%s（%s）" % (v["symbol"], v["meaning"], v["unit"]) for v in f["variables"]),
                    f["conditions"], None, None]
            for j, v in enumerate(vals, 1):
                cell = wf.cell(row=r, column=j, value=v)
                cell.font, cell.alignment, cell.border = f_body, wrap, border
                if j >= 8:
                    cell.fill = fill_input
            r += 1
flast = r - 1
wf.freeze_panes = "D2"
dv3 = DataValidation(type="list", formula1='"正确,有误,需讨论"', allow_blank=True)
wf.add_data_validation(dv3)
dv3.add("H2:H%d" % flast)

# ---------- 审校进度 ----------
wp = wb.create_sheet("审校进度")
heads = ["章节", "知识点数", "已审", "通过", "小改", "需重写", "完成率"]
for j, h in enumerate(heads, 1):
    c = wp.cell(row=1, column=j, value=h)
    c.font, c.fill, c.border = f_head, fill_head, border
    c.alignment = Alignment(horizontal="center")
rng_ch = "'知识点审校'!$B$2:$B$%d" % last
rng_q = "'知识点审校'!$Q$2:$Q$%d" % last
for i, ch in enumerate(course["chapters"], 2):
    name = "%d %s" % (ch["order"], ch["title"])
    wp.cell(row=i, column=1, value=name)
    wp.cell(row=i, column=2, value='=COUNTIF(%s,A%d)' % (rng_ch, i))
    wp.cell(row=i, column=3, value='=COUNTIFS(%s,A%d,%s,"<>")' % (rng_ch, i, rng_q))
    for j, k in zip((4, 5, 6), ("通过", "小改", "需重写")):
        wp.cell(row=i, column=j, value='=COUNTIFS(%s,A%d,%s,"%s")' % (rng_ch, i, rng_q, k))
    wp.cell(row=i, column=7, value="=IF(B%d=0,0,C%d/B%d)" % (i, i, i))
tot = len(course["chapters"]) + 2
wp.cell(row=tot, column=1, value="合计")
for j in range(2, 7):
    L = get_column_letter(j)
    wp.cell(row=tot, column=j, value="=SUM(%s2:%s%d)" % (L, L, tot - 1))
wp.cell(row=tot, column=7, value="=IF(B%d=0,0,C%d/B%d)" % (tot, tot, tot))
for row in wp.iter_rows(min_row=2, max_row=tot, max_col=7):
    for c in row:
        c.font, c.border = (f_bold if c.row == tot else f_body), border
    row[6].number_format = "0.0%"
wp.column_dimensions["A"].width = 22
for L in "BCDEFG":
    wp.column_dimensions[L].width = 11
wp.cell(row=tot + 2, column=1, value="本表全部为公式，按“知识点审校”表的审校结论自动统计。").font = Font(name=FONT, size=9, italic=True)

wb.save(out)
print("rows", last - 1, "formulas", flast - 1)
